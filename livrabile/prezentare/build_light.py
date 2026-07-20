#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import os

IMG = "/home/rciobanu/Theseis/prezentare/img"

# ---- Paleta Light (lizibila pe videoproiector) ----
BG     = RGBColor(0xFF, 0xFF, 0xFF)   # alb (fundal)
BG2    = RGBColor(0xEE, 0xF4, 0xF8)   # bleu foarte deschis (carduri)
CYAN   = RGBColor(0x0E, 0x7A, 0x99)   # teal-albastru profund (accent, lizibil pe alb)
CYAN_D = RGBColor(0x9C, 0xC5, 0xD4)   # teal deschis (linii carduri)
WHITE  = RGBColor(0x16, 0x27, 0x38)   # "text principal" = navy inchis (pe alb)
MUTE   = RGBColor(0x5A, 0x6B, 0x7B)   # text secundar gri-albastru
RED    = RGBColor(0xC0, 0x2A, 0x38)   # accent secundar (rar)
BAND   = RGBColor(0x11, 0x2A, 0x3D)   # navy pentru banda titlu slide 1

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background(); bg.shadow.inherit = False
    return s

def rect(s, x, y, w, h, color, line=None, lw=1.0):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh

def line(s, x, y, w, h, color):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    sh.line.fill.background(); sh.shadow.inherit = False
    return sh

def txt(s, x, y, w, h, text, size, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = "Segoe UI"
    return tb

def bullets(s, x, y, w, h, items, size=19, color=WHITE, gap=10, marker="▸ "):
    tb = s.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.05
        r1 = p.add_run(); r1.text = marker
        r1.font.size = Pt(size); r1.font.bold = True; r1.font.color.rgb = CYAN; r1.font.name = "Segoe UI"
        r2 = p.add_run(); r2.text = it
        r2.font.size = Pt(size); r2.font.color.rgb = color; r2.font.name = "Segoe UI"
    return tb

def header(s, kicker, title, n):
    # bara accent stanga sus
    line(s, Inches(0.55), Inches(0.55), Inches(0.12), Inches(0.55), CYAN)
    txt(s, Inches(0.8), Inches(0.5), Inches(9), Inches(0.35),
        kicker.upper(), 13, CYAN, bold=True)
    txt(s, Inches(0.8), Inches(0.82), Inches(11.8), Inches(0.85),
        title, 32, WHITE, bold=True)
    line(s, Inches(0.82), Inches(1.72), Inches(2.2), Pt(3), CYAN)
    # numar slide
    txt(s, SW-Inches(1.5), SH-Inches(0.55), Inches(1.0), Inches(0.35),
        f"{n:02d} / 14", 12, MUTE, align=PP_ALIGN.RIGHT)

def framed_img(s, path, x, y, maxw, maxh):
    """imagine incadrata cu rama cyan subtila, pastrand aspect ratio"""
    iw, ih = Image.open(path).size
    ar = iw/ih; box_ar = maxw/maxh
    if ar > box_ar: w = maxw; h = int(maxw/ar)
    else: h = maxh; w = int(maxh*ar)
    px = x + (maxw-w)//2; py = y + (maxh-h)//2
    # card fundal alb + rama teal
    pad = Inches(0.12)
    card = rect(s, px-pad, py-pad, w+2*pad, h+2*pad, RGBColor(0xFF,0xFF,0xFF))
    card.line.color.rgb = CYAN; card.line.width = Pt(1.25)
    s.shapes.add_picture(path, px, py, width=w, height=h)

def chip(s, x, y, w, h, title, body):
    """card mic cu titlu cyan + text"""
    c = rect(s, x, y, w, h, BG2, line=CYAN_D, lw=1.0)
    txt(s, x+Inches(0.2), y+Inches(0.15), w-Inches(0.4), Inches(0.4),
        title, 16, CYAN, bold=True)
    txt(s, x+Inches(0.2), y+Inches(0.62), w-Inches(0.4), h-Inches(0.7),
        body, 14, WHITE, spacing=1.05)

# ================= SLIDE 1 — TITLU =================
s = slide()
# accente geometrice
line(s, 0, Inches(2.55), SW, Pt(3), CYAN)
rect(s, Inches(0.9), Inches(2.85), Inches(0.9), Inches(0.12), CYAN)
txt(s, Inches(0.9), Inches(0.7), Inches(11), Inches(0.4),
    "UNIVERSITATEA TEHNICĂ DIN CLUJ-NAPOCA", 14, CYAN, bold=True)
txt(s, Inches(0.9), Inches(1.05), Inches(11), Inches(0.4),
    "Facultatea de Automatică și Calculatoare", 14, MUTE)
txt(s, Inches(0.9), Inches(3.05), Inches(11.5), Inches(2.0),
    "Sistem de Control FPGA pentru\nMână Robotică cu Filtrare Kalman",
    42, WHITE, bold=True, spacing=1.0)
txt(s, Inches(0.9), Inches(5.35), Inches(11), Inches(0.4),
    "LUCRARE DE LICENȚĂ", 18, CYAN, bold=True)
txt(s, Inches(0.9), Inches(6.05), Inches(11), Inches(0.4),
    "Absolvent:  Radu-Rareș CIOBANU", 18, WHITE)
txt(s, Inches(0.9), Inches(6.5), Inches(11), Inches(0.4),
    "Coordonator științific:  Prof. dr. ing. Radu Gabriel DĂNESCU   •   Iulie 2026", 14, MUTE)

# ================= SLIDE 2 — CONTEXT =================
s = slide()
header(s, "Introducere", "Context și motivație", 2)
bullets(s, Inches(0.8), Inches(2.1), Inches(6.7), Inches(4.6),
    ["Mâinile robotice care imită mâna umană sunt tot mai folosite pentru comandă la distanță, proteze și cercetare.",
     "Reproducerea mișcării umane cere control în timp real, precis și stabil.",
     "Unghiurile provin de la un sistem de viziune și conțin zgomot.",
     "Placa FPGA a fost aleasă pentru viteza și paralelismul ei, iar filtrul Kalman face mișcarea lină."],
    size=19, gap=16)
framed_img(s, f"{IMG}/mana_completa.jpg", Inches(7.8), Inches(2.1), Inches(4.9), Inches(4.6))

# ================= SLIDE 3 — OBIECTIVE =================
s = slide()
header(s, "Obiective", "Obiectivele proiectului", 3)
bullets(s, Inches(1.0), Inches(2.2), Inches(11.3), Inches(4.6),
    ["Controlul în timp real al unei mâini robotice cu 8 servomotoare.",
     "Recepția unghiurilor prin UART și filtrarea lor cu un filtru Kalman.",
     "Generarea celor 8 semnale PWM de comandă către servomotoare.",
     "Citirea poziției reale a servomotoarelor printr-un convertor analog-digital.",
     "Izolarea galvanică între partea logică și partea de putere.",
     "Implementare integrală în virgulă fixă, pe un circuit de 100 MHz."],
    size=20, gap=13)

# ================= SLIDE 4 — ARHITECTURA =================
s = slide()
header(s, "Privire de ansamblu", "Arhitectura sistemului", 4)
framed_img(s, f"{IMG}/arhitectura_module.png", Inches(0.8), Inches(2.0), Inches(8.3), Inches(4.9))
# panou lateral cu fluxul
chip(s, Inches(9.5), Inches(2.05), Inches(3.3), Inches(1.35),
     "Comandă", "UART → cutie poștală → filtru Kalman → generare PWM")
chip(s, Inches(9.5), Inches(3.55), Inches(3.3), Inches(1.35),
     "Feedback", "ADC AD7124-8 citit prin SPI, izolat cu MAX14850")
chip(s, Inches(9.5), Inches(5.05), Inches(3.3), Inches(1.35),
     "Izolare", "ADuM1400 separă comanda de zgomotul servomotoarelor")

# ================= SLIDE 5 — BASYS3 =================
s = slide()
header(s, "Componentă cheie · 1", "Placa FPGA Basys3", 5)
bullets(s, Inches(0.8), Inches(2.1), Inches(6.6), Inches(4.6),
    ["Circuit Xilinx Artix-7 XC7A35T, ceas de 100 MHz.",
     "Găzduiește toate cele 14 module VHDL ale sistemului.",
     "Rulează în paralel: UART, Kalman, PWM și bucla SPI.",
     "Determinism total și închidere temporală respectată.",
     "Moduri de lucru: normal, demonstrativ și manual."],
    size=19, gap=15)
framed_img(s, f"{IMG}/arhitectura_module.png", Inches(7.7), Inches(2.3), Inches(5.0), Inches(4.0))

# ================= SLIDE 6 — KALMAN =================
s = slide()
header(s, "Contribuția centrală", "Filtrul Kalman", 6)
bullets(s, Inches(0.8), Inches(2.1), Inches(6.4), Inches(4.6),
    ["Filtru cu două stări: poziție și viteză.",
     "Implementat integral în aritmetică cu virgulă fixă Q16.16.",
     "Prelucrează toate cele 8 canale într-un ciclu de control.",
     "Împărțirea realizată pe mai multe cicluri (restoring divider).",
     "Filtrează zgomotul din unghiuri și dă o mișcare lină."],
    size=19, gap=15)
framed_img(s, f"{IMG}/kalman_fsm.png", Inches(7.7), Inches(1.95), Inches(5.0), Inches(5.0))

# ================= SLIDE 7 — AD7124 =================
s = slide()
header(s, "Componentă cheie · 2", "Convertorul AD7124-8", 7)
bullets(s, Inches(0.8), Inches(2.1), Inches(6.5), Inches(4.6),
    ["Convertor sigma-delta pe 24 de biți (Analog Devices).",
     "Citește poziția reală a servomotoarelor — feedback.",
     "Comunicație SPI, 8 canale, prin divizoare rezistive.",
     "Referință internă și buffere de intrare activate.",
     "Închide bucla: comandă → mișcare → măsurare."],
    size=19, gap=15)
framed_img(s, f"{IMG}/adi_ad7124.png", Inches(7.7), Inches(2.3), Inches(5.0), Inches(4.1))

# ================= SLIDE 8 — IZOLARE =================
s = slide()
header(s, "Componentă cheie · 3", "Izolarea galvanică", 8)
bullets(s, Inches(0.8), Inches(2.05), Inches(6.4), Inches(4.7),
    ["Separă partea logică de zgomotul servomotoarelor.",
     "ADuM1400 (×2): izolează cele 8 semnale PWM.",
     "Tehnologie iCoupler, cu transformatoare integrate.",
     "MAX14850: izolează magistrala SPI către convertor.",
     "Previne propagarea vârfurilor de tensiune în placă."],
    size=18, gap=13)
framed_img(s, f"{IMG}/adi_adum_diagrama.png", Inches(7.6), Inches(2.1), Inches(5.1), Inches(4.6))

# ================= SLIDE 9 — REALIZARE FIZICA =================
s = slide()
header(s, "Ansamblul mecanic", "Realizarea fizică", 9)
bullets(s, Inches(0.8), Inches(2.2), Inches(5.0), Inches(4.4),
    ["Mână printată 3D, ansamblu hibrid cu piese proprii.",
     "Acționare prin fire și arcuri (tendoane).",
     "8 servomotoare MG996R.",
     "Feedback prin intervenție directă pe servomotoare.",
     "Integrare cu placa FPGA și alimentarea."],
    size=18, gap=13)
framed_img(s, f"{IMG}/brat_real.jpg", Inches(6.1), Inches(1.95), Inches(6.6), Inches(5.0))

# ================= SLIDE 10 — TESTARE & REZULTATE =================
s = slide()
header(s, "Validare", "Testare și rezultate", 10)
bullets(s, Inches(0.8), Inches(2.1), Inches(6.4), Inches(4.6),
    ["Testare pe niveluri: fiecare modul verificat separat.",
     "Închidere temporală confirmată în Vivado la 100 MHz.",
     "Comunicație SPI verificată cu osciloscopul.",
     "Toate cele 8 canale ale ADC citite corect.",
     "Mâna reproduce mișcările și prinde obiecte reale."],
    size=19, gap=15)
framed_img(s, f"{IMG}/timing_pwm.png", Inches(7.7), Inches(2.4), Inches(5.0), Inches(2.9))
txt(s, Inches(7.7), Inches(5.6), Inches(5.0), Inches(0.4),
    "Semnal PWM: 0,544 – 2,4 ms", 13, MUTE, align=PP_ALIGN.CENTER, italic=True)

# ================= SLIDE 11 — CONTRIBUTII =================
s = slide()
header(s, "Aportul autorului", "Contribuții personale", 11)
# 3 carduri late
data = [
    ("Filtrul Kalman", "Filtru cu două stări în virgulă fixă Q16.16, implementat pe FPGA, adaptat unui circuit care lucrează doar cu numere întregi."),
    ("Lanțul complet", "De la recepția UART, prin filtrare și sincronizare, până la generarea PWM și feedback-ul prin SPI."),
    ("Integrarea fizică", "Controler dedicat pentru AD7124-8, realizarea mâinii și integrarea componentelor Analog Devices."),
]
y = Inches(2.15)
for i,(t,b) in enumerate(data):
    chip(s, Inches(0.8), y, Inches(11.7), Inches(1.35), t, b)
    y += Inches(1.55)

# ================= SLIDE 12 — CONCLUZII =================
s = slide()
header(s, "Final", "Concluzii și direcții viitoare", 12)
txt(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.6),
    "Sistemul funcționează stabil, iar mâna reproduce mișcările și prinde obiecte reale.",
    20, WHITE, bold=True)
txt(s, Inches(0.8), Inches(2.9), Inches(6), Inches(0.4), "DIRECȚII VIITOARE", 15, CYAN, bold=True)
bullets(s, Inches(0.8), Inches(3.35), Inches(11.5), Inches(2.8),
    ["Material antiderapant la vârful degetelor pentru prindere mai sigură.",
     "Alimentare la 6 V pentru cuplu și forță de strângere mai mari.",
     "Comunicație fără fir în locul legăturii seriale.",
     "Filtru Kalman extins la un model cu trei stări.",
     "Integrarea completă cu partea de viziune pentru teleoperare."],
    size=18, gap=11)

# ================= SLIDE 13 — REFERINTE =================
s = slide()
header(s, "Bibliografie", "Referințe bibliografice", 13)
refs = [
    "R. E. Kalman, „A New Approach to Linear Filtering and Prediction Problems”, Journal of Basic Engineering, 1960.",
    "Analog Devices, „AD7124-8 — Low Power, Low Noise, 24-Bit Sigma-Delta ADC”, fișă tehnică.",
    "Analog Devices, „ADuM1400 Quad-Channel Digital Isolator” și „MAX14850 6-Channel Digital Isolator”, fișe tehnice.",
    "Analog Devices, „ADP165 Low Noise LDO Regulator” și „DC2468A Power Module”, fișe tehnice.",
    "Google, „MediaPipe Hands: On-device Real-time Hand Tracking”, 2020.",
    "Digilent, „Basys3 FPGA Board Reference Manual”.",
    "Proiectul InMoov — mână robotică open-source (inmoov.fr).",
]
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.05), Inches(11.9), Inches(5.0))
tf = tb.text_frame; tf.word_wrap = True
for i, r in enumerate(refs):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(10); p.line_spacing = 1.05
    r1 = p.add_run(); r1.text = f"[{i+1}]  "
    r1.font.size = Pt(15); r1.font.bold = True; r1.font.color.rgb = CYAN; r1.font.name = "Segoe UI"
    r2 = p.add_run(); r2.text = r
    r2.font.size = Pt(15); r2.font.color.rgb = WHITE; r2.font.name = "Segoe UI"

# ================= SLIDE 14 — MULTUMESC =================
s = slide()
# banda accent
line(s, 0, Inches(3.1), SW, Pt(4), CYAN)
rect(s, Inches(5.5), Inches(3.4), Inches(2.3), Inches(0.14), CYAN)
txt(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.0),
    "Vă mulțumesc pentru atenție!", 40, WHITE, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(3.75), Inches(11.7), Inches(0.5),
    "Sistem de Control FPGA pentru Mână Robotică cu Filtrare Kalman", 18, CYAN, bold=True, align=PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.4),
    "Radu-Rareș CIOBANU", 18, WHITE, align=PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(4.95), Inches(11.7), Inches(0.4),
    "Coordonator: Prof. dr. ing. Radu Gabriel DĂNESCU", 14, MUTE, align=PP_ALIGN.CENTER)
txt(s, Inches(0.8), Inches(5.4), Inches(11.7), Inches(0.4),
    "Universitatea Tehnică din Cluj-Napoca  •  Iulie 2026", 14, MUTE, align=PP_ALIGN.CENTER)

out = "/tmp/ppt/prezentare_dark.pptx"
prs.save(out)
print("SALVAT:", out, "-", len(prs.slides._sldIdLst), "slide-uri")
